# -*- coding: utf-8 -*-
import io

from openpyxl import load_workbook

from odoo import http
from odoo.tests import HttpCase, tagged
from odoo.tools import mute_logger

from .common import DakoPasswordData

ACTION_URL = '/odoo/action-dako_password.dako_password_action'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


@tagged('post_install', '-at_install')
class TestDakoPasswordExportRoute(HttpCase, DakoPasswordData):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._dako_setup_data()

    def _post_export(self, **data):
        return self.url_open('/dako_password/export/xlsx', data={
            'domain': '[]', 'ids': '[]', 'csrf_token': http.Request.csrf_token(self), **data,
        })

    def test_export_route(self):
        self.authenticate('dako_a', 'dako_axx')
        response = self._post_export()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['Content-Type'], XLSX)
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual([row[1] for row in rows[1:]], ['Banco A'])
        self.assertEqual(len(self._logs(self.cred_a, 'export')), 1)

    def test_export_route_selected_ids(self):
        self.authenticate('dako_m', 'dako_mxx')
        response = self._post_export(ids='[%d]' % self.cred_b.id)
        rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual([row[1] for row in rows[1:]], ['Servidor B'])

    @mute_logger('odoo.http')
    def test_export_route_refuses_get_and_users_without_export_right(self):
        self.authenticate('dako_a', 'dako_axx')
        self.assertEqual(self.url_open('/dako_password/export/xlsx').status_code, 405)
        self.authenticate('dako_b', 'dako_bxx')
        response = self._post_export()
        self.assertEqual(response.status_code, 500)
        self.assertIn('odoo.exceptions.AccessError', response.text)
        self.assertFalse(self._logs(self.cred_b))

    @mute_logger('odoo.http')
    def test_export_route_refuses_a_domain_on_the_password(self):
        self.authenticate('dako_a', 'dako_axx')
        response = self._post_export(domain='[["password", "=like", "S%"]]')
        self.assertEqual(response.status_code, 500)
        self.assertFalse(self._logs(self.cred_a))


@tagged('post_install', '-at_install')
class TestDakoPasswordTour(HttpCase, DakoPasswordData):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._dako_setup_data()

    def _check_desktop_tour(self):
        self.assertEqual(self.cred_a.sudo().password, 'Cambiada-Tour-2026!')
        self.assertEqual(self.cred_a.strength, 'strong')
        # Eye, then "Mostrar todas": two reveals. Copy accepted, then refused: two copies.
        self.assertEqual(len(self._logs(self.cred_a, 'reveal')), 2)
        self.assertEqual(len(self._logs(self.cred_a, 'copy')), 2)
        self.assertFalse(self._logs(self.cred_b))

    def test_desktop_light(self):
        self.start_tour(ACTION_URL, 'dako_password_desktop', login='dako_a')
        self._check_desktop_tour()

    def test_desktop_dark(self):
        self.start_tour(ACTION_URL, 'dako_password_desktop', login='dako_a', cookies={'color_scheme': 'dark'})
        self._check_desktop_tour()


@tagged('post_install', '-at_install')
class TestDakoPasswordMobileTour(HttpCase, DakoPasswordData):
    browser_size = '375x667'
    touch_enabled = True

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._dako_setup_data()

    def test_mobile(self):
        self.start_tour(ACTION_URL, 'dako_password_mobile', login='dako_a')
        self.assertTrue(self._logs(self.cred_a, 'reveal'))
        self.assertEqual(len(self._logs(self.cred_a, 'copy')), 1)

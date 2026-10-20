# -*- coding: utf-8 -*-
import io
from datetime import datetime

from openpyxl import load_workbook

from odoo import Command
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests import Form, tagged
from odoo.tools import mute_logger

from odoo.addons.dako_password.models.dako_password import (
    DIGITS,
    LOWERCASE,
    SYMBOLS,
    UPPERCASE,
    new_password,
    password_strength,
)

from .common import DakoPasswordCommon


@tagged('post_install', '-at_install')
class TestDakoPasswordAccess(DakoPasswordCommon):

    def test_member_sees_only_own_vaults(self):
        Vault = self.env['dako.password.vault'].with_user(self.user_a)
        Password = self.env['dako.password'].with_user(self.user_a)
        self.assertEqual(Vault.search([]), self.vault_a)
        self.assertEqual(Password.search([]), self.cred_a)

    def test_manager_sees_everything(self):
        # The database may hold other vaults: the manager must see at least these ones,
        # without being a member of any of them.
        Vault = self.env['dako.password.vault'].with_user(self.manager)
        Password = self.env['dako.password'].with_user(self.manager)
        self.assertLessEqual(self.vault_a | self.vault_b, Vault.search([]))
        self.assertLessEqual(self.cred_a | self.cred_b, Password.search([]))

    @mute_logger('odoo.models', 'odoo.addons.base.models.ir_rule')
    def test_member_cannot_create_in_foreign_vault(self):
        # _create() validates constraints before the record rules: the clearer message wins.
        with self.assertRaises(ValidationError):
            self.env['dako.password'].with_user(self.user_a).create({
                'name': 'Intruso', 'password_input': 'x', 'vault_id': self.vault_b.id,
            })

    @mute_logger('odoo.models')
    def test_only_manager_deletes_credentials(self):
        with self.assertRaises(AccessError):
            self.cred_a.with_user(self.user_a).unlink()
        self.cred_a.with_user(self.manager).unlink()
        self.assertFalse(self.cred_a.exists())

    @mute_logger('odoo.models')
    def test_access_log_is_read_only_for_manager_and_hidden_for_members(self):
        self.cred_a.with_user(self.user_a).dako_reveal('reveal')
        with self.assertRaises(AccessError):
            self.Log.with_user(self.user_a).search([])
        logs = self.Log.with_user(self.manager).search([])
        self.assertTrue(logs)
        with self.assertRaises(AccessError):
            logs.with_user(self.manager).write({'action': 'copy'})
        with self.assertRaises(AccessError):
            logs.with_user(self.manager).unlink()

    def test_creator_is_added_to_vault_unless_manager(self):
        Vault = self.env['dako.password.vault']
        own = Vault.with_user(self.user_a).create({
            'name': 'Mio', 'user_ids': [Command.set(self.user_b.ids)],
        })
        self.assertEqual(own.user_ids, self.user_a | self.user_b)
        managed = Vault.with_user(self.manager).create({
            'name': 'Para B', 'user_ids': [Command.set(self.user_b.ids)],
        })
        self.assertEqual(managed.user_ids, self.user_b)

    def test_role_selector_only_in_debug_mode(self):
        # Odoo 19: a group without privilege goes to the "Extra Rights" section of the
        # user form, which the web client only renders in debug (technical) mode.
        group = self.env.ref('dako_password.dako_password_group_manager')
        self.assertFalse(group.privilege_id)
        hierarchy = self.env['res.groups']._get_view_group_hierarchy()
        self.assertFalse(hierarchy['groups'][group.id]['privilege_id'])

    def test_default_vault_when_user_has_only_one(self):
        Password = self.env['dako.password'].with_user(self.user_a)
        self.assertEqual(Password.default_get(['vault_id'])['vault_id'], self.vault_a.id)
        self.vault_b.user_ids = [Command.link(self.user_a.id)]
        self.assertFalse(Password.default_get(['vault_id']).get('vault_id'))


@tagged('post_install', '-at_install')
class TestDakoPasswordSecret(DakoPasswordCommon):

    def test_regular_read_does_not_return_password(self):
        # Every field of this module the user can see, as the web client lists them. A bare
        # read() would also fetch fields that other modules add to every mail.thread model
        # (e.g. OCA mgmtsystem_nonconformity) and that this user may not be allowed to read.
        cred = self.cred_a.with_user(self.user_a)
        names = [name for name in cred.fields_get() if 'dako_password' in cred._fields[name]._modules]
        self.assertIn('password_input', names)
        values = cred.read(names)[0]
        self.assertNotIn('password', values)
        self.assertFalse(values['password_input'])
        self.assertNotIn('Secreto-A-2026', str(values))

    @mute_logger('odoo.models', 'odoo.fields')
    def test_every_regular_path_refuses_password(self):
        cred = self.cred_a.with_user(self.user_a)
        Password = self.env['dako.password'].with_user(self.user_a)
        calls = {
            'read': lambda: cred.read(['password']),
            'web_read': lambda: cred.web_read({'password': {}}),
            'search_read': lambda: Password.search_read([], ['password']),
            'search': lambda: Password.search([('password', '=like', 'S%')]),
            'order': lambda: Password.search([], order='password'),
            'read_group': lambda: Password._read_group([], ['password'], ['__count']),
            'export_data': lambda: cred.export_data(['password']),
            'onchange': lambda: cred.onchange({}, [], {'password': {}}),
            'write': lambda: cred.write({'password': 'Robado-1'}),
        }
        for name, call in calls.items():
            with self.subTest(path=name), self.assertRaises(AccessError):
                call()
        self.assertEqual(self.cred_a.sudo().password, 'Secreto-A-2026')

    def test_manager_cannot_read_password_either(self):
        with self.assertRaises(AccessError), mute_logger('odoo.models'):
            self.cred_a.with_user(self.manager).read(['password'])

    def test_fields_get_and_views_hide_password(self):
        Password = self.env['dako.password'].with_user(self.user_a)
        self.assertNotIn('password', Password.fields_get())
        views = Password.get_views([(False, 'form'), (False, 'list'), (False, 'kanban'), (False, 'search')])
        for view_type, view in views['views'].items():
            with self.subTest(view=view_type):
                self.assertNotIn('name="password"', view['arch'])

    def test_change_password_through_input(self):
        self.cred_a.sudo().password_changed_date = datetime(2020, 1, 1)
        self._flush_tracking()
        messages_before = self.cred_a.message_ids
        self.cred_a.with_user(self.user_a).write({'password_input': 'Nuevo-Valor-99'})
        self._flush_tracking()
        self.assertEqual(self.cred_a.sudo().password, 'Nuevo-Valor-99')
        self.assertGreater(self.cred_a.password_changed_date, datetime(2020, 1, 1))
        new_messages = self.cred_a.message_ids - messages_before
        self.assertIn('Cambió la contraseña.', new_messages.mapped('body')[0])
        self.assertEqual(new_messages.author_id, self.user_a.partner_id)
        for message in new_messages:
            self.assertNotIn('Nuevo-Valor-99', message.body or '')
            for tracking in message.sudo().tracking_value_ids:
                self.assertNotIn('Nuevo-Valor-99', '%s %s' % (tracking.old_value_char, tracking.new_value_char))

    def test_same_password_is_not_a_change(self):
        self._flush_tracking()
        messages_before = self.cred_a.message_ids
        self.cred_a.with_user(self.user_a).write({'password_input': 'Secreto-A-2026'})
        self._flush_tracking()
        self.assertFalse(self.cred_a.message_ids - messages_before)

    def test_create_requires_password(self):
        with self.assertRaises(ValidationError):
            self.env['dako.password'].with_user(self.user_a).create({
                'name': 'Sin clave', 'vault_id': self.vault_a.id,
            })

    def test_create_with_input_computes_strength(self):
        cred = self.env['dako.password'].with_user(self.user_a).create({
            'name': 'Nueva', 'vault_id': self.vault_a.id, 'password_input': 'abc',
        })
        self.assertEqual(cred.sudo().password, 'abc')
        self.assertEqual(cred.strength, 'weak')

    def test_notes_change_is_logged_without_content(self):
        self._flush_tracking()
        messages_before = self.cred_a.message_ids
        self.cred_a.with_user(self.user_a).write({'note': 'PIN 4321'})
        self._flush_tracking()
        bodies = ' '.join((self.cred_a.message_ids - messages_before).mapped('body'))
        self.assertIn('Modificó las notas.', bodies)
        self.assertNotIn('4321', bodies)

    def test_duplicate_keeps_password_and_logs(self):
        copy = self.cred_a.with_user(self.user_a).copy()
        self.assertEqual(copy.name, 'Banco A (copia)')
        self.assertEqual(copy.sudo().password, 'Secreto-A-2026')
        self.assertEqual(copy.strength, self.cred_a.strength)
        self.assertEqual(len(self._logs(self.cred_a, 'duplicate')), 1)
        self.assertNotIn('password', self.cred_a.with_user(self.user_a).copy_data()[0])


@tagged('post_install', '-at_install')
class TestDakoPasswordReveal(DakoPasswordCommon):

    def test_reveal_returns_value_and_logs(self):
        result = self.cred_a.with_user(self.user_a).dako_reveal('reveal')
        self.assertEqual(result, {self.cred_a.id: 'Secreto-A-2026'})
        log = self._logs(self.cred_a)
        self.assertEqual(len(log), 1)
        self.assertEqual(log.user_id, self.user_a)
        self.assertEqual(log.action, 'reveal')
        self.assertEqual(log.vault_id, self.vault_a)
        self.assertEqual(log.credential_name, 'Banco A')

    @mute_logger('odoo.models', 'odoo.addons.base.models.ir_rule')
    def test_reveal_foreign_credential_is_refused_and_not_logged(self):
        with self.assertRaises(AccessError):
            self.cred_b.with_user(self.user_a).dako_reveal('copy')
        self.assertFalse(self._logs(self.cred_b))

    def test_unknown_action_is_refused(self):
        with self.assertRaises(ValidationError):
            self.cred_a.with_user(self.user_a).dako_reveal('export')

    def test_reveal_batch_logs_each_credential(self):
        both = (self.cred_a | self.cred_b).with_user(self.manager)
        result = both.dako_reveal('reveal')
        self.assertEqual(set(result), {self.cred_a.id, self.cred_b.id})
        self.assertEqual(len(self.Log.search([('user_id', '=', self.manager.id)])), 2)

    def test_reveal_and_copy_stay_out_of_the_chatter(self):
        self._flush_tracking()
        messages_before = self.cred_a.message_ids
        self.cred_a.with_user(self.user_a).dako_reveal('reveal')
        self.cred_a.with_user(self.user_a).dako_reveal('copy')
        self._flush_tracking()
        self.assertFalse(self.cred_a.message_ids - messages_before)
        self.assertEqual(len(self._logs(self.cred_a)), 2)

    def test_access_log_count_is_manager_only(self):
        self.cred_a.with_user(self.user_a).dako_reveal('reveal')
        self.assertEqual(self.cred_a.with_user(self.manager).access_log_count, 1)
        with self.assertRaises(AccessError), mute_logger('odoo.models'):
            self.cred_a.with_user(self.user_a).read(['access_log_count'])


@tagged('post_install', '-at_install')
class TestDakoPasswordStrength(DakoPasswordCommon):

    def test_strength_table(self):
        cases = [
            ('', (0, 'weak')),
            ('abc', (0, 'weak')),
            ('abcdefgh', (25, 'weak')),
            ('abcdefgh1', (45, 'medium')),
            ('Abcdefgh1', (65, 'medium')),
            ('Abcdefgh1!', (85, 'strong')),
            ('Abcdefghijk1!', (100, 'strong')),
        ]
        for value, expected in cases:
            with self.subTest(value=value):
                self.assertEqual(password_strength(value), expected)

    def test_generator(self):
        allowed = set(LOWERCASE + UPPERCASE + DIGITS + SYMBOLS)
        values = set()
        for _i in range(200):
            value = new_password()
            values.add(value)
            self.assertEqual(len(value), 20)
            self.assertTrue(set(value) <= allowed)
            self.assertFalse(set(value) & set('O0Il1'))
            for chars_class in (LOWERCASE, UPPERCASE, DIGITS, SYMBOLS):
                self.assertTrue(set(value) & set(chars_class))
            self.assertEqual(password_strength(value)[1], 'strong')
        self.assertEqual(len(values), 200)
        self.assertEqual(len(self.env['dako.password'].with_user(self.user_a).generate_password()), 20)

    def test_form_shows_strength_while_typing(self):
        with Form(self.cred_a.with_user(self.user_a)) as form:
            form.password_input = 'abc'
            self.assertEqual(form.strength, 'weak')
            form.password_input = 'Muy-Segura-2026!'
            self.assertEqual(form.strength, 'strong')
            self.assertEqual(form.strength_score, 100)
        self.assertEqual(self.cred_a.sudo().password, 'Muy-Segura-2026!')
        self.assertEqual(self.cred_a.strength, 'strong')

    def test_form_creates_credential(self):
        with Form(self.env['dako.password'].with_user(self.user_a)) as form:
            form.name = 'Desde la ficha'
            form.password_input = 'Clave-De-Ficha-1'
        cred = form.record
        self.assertEqual(cred.vault_id, self.vault_a)
        self.assertEqual(cred.sudo().password, 'Clave-De-Ficha-1')
        self.assertEqual(cred.criticality, 'medium')

    def test_url_with_credentials_warns(self):
        cred = self.env['dako.password'].new({'url': 'https://admin:clave@intranet.example.com'})
        self.assertIn('warning', cred._onchange_url() or {})
        cred.url = '192.168.1.10:3389'
        self.assertFalse(cred._onchange_url())


@tagged('post_install', '-at_install')
class TestDakoPasswordVault(DakoPasswordCommon):

    def test_vault_needs_an_internal_user(self):
        with self.assertRaises(ValidationError):
            self.vault_a.write({'user_ids': [Command.clear()]})
        portal = self.env['res.users'].create({
            'name': 'Portal', 'login': 'dako_portal',
            'group_ids': [Command.set(self.env.ref('base.group_portal').ids)],
        })
        with self.assertRaises(ValidationError):
            self.vault_a.write({'user_ids': [Command.link(portal.id)]})

    def test_member_cannot_leave_but_manager_can_remove(self):
        self.vault_a.user_ids = [Command.link(self.user_b.id)]
        with self.assertRaises(UserError):
            self.vault_a.with_user(self.user_a).write({'user_ids': [Command.unlink(self.user_a.id)]})
        self.vault_a.with_user(self.user_a).write({'user_ids': [Command.unlink(self.user_b.id)]})
        self.vault_a.with_user(self.manager).write({'user_ids': [Command.set(self.user_b.ids)]})
        self.assertEqual(self.vault_a.user_ids, self.user_b)

    @mute_logger('odoo.models', 'odoo.addons.base.models.ir_rule')
    def test_cannot_move_credential_to_foreign_vault(self):
        with self.assertRaises(ValidationError):
            self.cred_a.with_user(self.user_a).write({'vault_id': self.vault_b.id})
        self.cred_a.with_user(self.manager).write({'vault_id': self.vault_b.id})
        self.assertEqual(self.cred_a.vault_id, self.vault_b)

    def test_archiving_vault_archives_credentials(self):
        retired = self.env['dako.password'].create({
            'name': 'Dada de baja', 'password_input': 'x', 'vault_id': self.vault_a.id,
        })
        retired.with_user(self.user_a).action_archive()
        self.vault_a.with_user(self.user_a).action_archive()
        self.assertFalse(self.cred_a.active)
        self.vault_a.with_user(self.user_a).action_unarchive()
        self.assertTrue(self.cred_a.active)
        # Archived on purpose before the vault: it stays archived.
        self.assertFalse(retired.active)

    @mute_logger('odoo.models')
    def test_vault_with_credentials_cannot_be_deleted(self):
        with self.assertRaises(UserError):
            self.vault_a.with_user(self.manager).unlink()
        empty = self.env['dako.password.vault'].create({
            'name': 'Vacio', 'user_ids': [Command.set(self.user_a.ids)],
        })
        with self.assertRaises(AccessError):
            empty.with_user(self.user_a).unlink()
        empty.with_user(self.manager).unlink()
        self.assertFalse(empty.exists())

    def test_user_changes_are_tracked(self):
        self._flush_tracking()
        messages_before = self.vault_a.message_ids
        self.vault_a.with_user(self.user_a).write({'user_ids': [Command.link(self.user_b.id)]})
        self._flush_tracking()
        tracking = (self.vault_a.message_ids - messages_before).sudo().tracking_value_ids
        self.assertIn('user_ids', tracking.field_id.mapped('name'))

    def test_counters(self):
        self.env['dako.password'].create({
            'name': 'Otra', 'password_input': 'x', 'vault_id': self.vault_a.id, 'criticality': 'low',
        })
        vault = self.vault_a.with_user(self.user_a)
        self.assertEqual(vault.password_count, 2)
        self.assertEqual(vault.critical_count, 1)

    def test_open_credentials_from_vault(self):
        action = self.vault_a.with_user(self.user_a).action_open_passwords()
        self.assertEqual(action['res_model'], 'dako.password')
        self.assertEqual(action['context']['default_vault_id'], self.vault_a.id)
        self.assertEqual(action['context']['searchpanel_default_vault_id'], self.vault_a.id)


@tagged('post_install', '-at_install')
class TestDakoPasswordExport(DakoPasswordCommon):

    def _rows(self, content):
        sheet = load_workbook(io.BytesIO(content)).active
        return list(sheet.iter_rows(values_only=True)), sheet

    def test_export_contains_only_accessible_credentials_and_logs(self):
        records = self.env['dako.password'].with_user(self.user_a).search([])
        rows, _sheet = self._rows(records._dako_export_xlsx())
        self.assertEqual(len(rows), 2)
        header, row = rows
        self.assertEqual(header[5], 'Contraseña')
        self.assertEqual(row[:7], (
            'Vault A', 'Banco A', 'Pagos a proveedores', 'https://banco.example.com',
            'tesoreria@example.com', 'Secreto-A-2026', 'Crítica',
        ))
        self.assertEqual(len(self._logs(self.cred_a, 'export')), 1)
        self.assertFalse(self._logs(self.cred_b))

    def test_export_keeps_formulas_as_text(self):
        formula = '=HYPERLINK("http://example.com","x")'
        self.cred_a.write({'name': formula, 'password_input': '=1+1'})
        rows, sheet = self._rows(self.cred_a.with_user(self.user_a)._dako_export_xlsx())
        self.assertEqual(rows[1][1], formula)
        self.assertEqual(rows[1][5], '=1+1')
        self.assertEqual(sheet.cell(row=2, column=2).data_type, 's')
        self.assertEqual(sheet.cell(row=2, column=6).data_type, 's')
